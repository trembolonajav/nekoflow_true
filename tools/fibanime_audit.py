#!/usr/bin/env python3
"""Coleta auditavel do catalogo publico do FIB Anime.

O coletor le as URLs da aba ``Animes`` da planilha de inventario, percorre
uma obra por vez e salva checkpoints apos cada anime. Ele registra apenas os
metadados e as fontes declaradas no ``textarea#player-data``; nao baixa midia.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

from lxml import html as lxml_html
from openpyxl import load_workbook


BASE_URL = "https://fibanime.cc"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/155 Safari/537.36 "
    "NekoflowCatalogAudit/1.0"
)
EPISODE_RE = re.compile(r"T\s*(\d+)\s*EP\s*(\d+)", re.IGNORECASE)


@dataclass
class FetchResult:
    url: str
    status: int
    text: str


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_text(parts: Iterable[str]) -> str:
    return " ".join(" ".join(parts).split())


def fetch(url: str, *, retries: int = 3, timeout: int = 35) -> FetchResult:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            req = Request(
                url,
                headers={
                    "User-Agent": USER_AGENT,
                    "Accept": "text/html,application/xhtml+xml",
                    "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.6",
                    "Cache-Control": "no-cache",
                },
            )
            with urlopen(req, timeout=timeout) as response:
                raw = response.read()
                charset = response.headers.get_content_charset() or "utf-8"
                return FetchResult(
                    url=response.geturl(),
                    status=response.status,
                    text=raw.decode(charset, errors="replace"),
                )
        except (HTTPError, URLError, TimeoutError, OSError) as exc:
            last_error = exc
            if isinstance(exc, HTTPError) and exc.code in {400, 401, 403, 404}:
                raise
            if attempt + 1 < retries:
                time.sleep((2**attempt) + random.uniform(0.1, 0.5))
    assert last_error is not None
    raise last_error


def load_inventory(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Animes" not in workbook.sheetnames:
        raise ValueError("A planilha nao possui a aba 'Animes'.")
    sheet = workbook["Animes"]
    records: list[dict[str, Any]] = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        url_index = next(
            (
                index
                for index, value in enumerate(values)
                if isinstance(value, str) and "/anime/" in value
            ),
            None,
        )
        if url_index is None:
            continue
        url = str(values[url_index]).strip().split("?")[0]
        if not url.startswith("http"):
            continue
        order = values[0] if isinstance(values[0], int) else len(records) + 1
        title = ""
        if url_index >= 2 and isinstance(values[1], str):
            title = values[1].strip()
        slug = urlparse(url).path.rstrip("/").split("/")[-1]
        records.append({"order": int(order), "title": title or slug, "slug": slug, "url": url})
    unique = {record["url"]: record for record in records}
    return sorted(unique.values(), key=lambda item: item["order"])


def parse_anime_page(page_url: str, document: str) -> dict[str, Any]:
    root = lxml_html.fromstring(document)
    root.make_links_absolute(page_url)
    heading = root.xpath("//main//h1[1]//text()")
    title = clean_text(heading) if heading else ""
    episodes_by_key: dict[tuple[int, int], dict[str, Any]] = {}
    season_meta: dict[int, dict[str, Any]] = {}

    for anchor in root.xpath('//a[contains(@href, "/assistir/")]'):
        text = clean_text(anchor.xpath(".//text()"))
        match = EPISODE_RE.search(text)
        if not match:
            continue
        season_number, episode_number = map(int, match.groups())
        episode_url = anchor.get("href") or ""
        key = (season_number, episode_number)
        episodes_by_key[key] = {
            "season": season_number,
            "episode": episode_number,
            "code": f"T{season_number} EP{episode_number:02d}",
            "url": urljoin(page_url, episode_url),
            "status": "pending",
            "sources": [],
        }

        article = anchor.xpath("ancestor::article[1]")
        if article and season_number not in season_meta:
            article_node = article[0]
            article_text = clean_text(article_node.xpath(".//text()"))
            expected_match = re.search(r"Epis[oó]dios:\s*(\d+)", article_text, re.IGNORECASE)
            custom_names = article_node.xpath(
                './/*[contains(concat(" ", normalize-space(@class), " "), " season-custom-name ")]//text()'
            )
            panel_names = article_node.xpath(
                './/*[contains(concat(" ", normalize-space(@class), " "), " season-panel-name ")]//text()'
            )
            season_title = clean_text(custom_names or panel_names) or title
            season_meta[season_number] = {
                "season": season_number,
                "title": season_title,
                "expected_episodes": int(expected_match.group(1)) if expected_match else None,
            }

    episodes = [episodes_by_key[key] for key in sorted(episodes_by_key)]
    for season_number in sorted({episode["season"] for episode in episodes}):
        if season_number not in season_meta:
            season_meta[season_number] = {
                "season": season_number,
                "title": "",
                "expected_episodes": None,
            }
        season_meta[season_number]["listed_episodes"] = sum(
            1 for episode in episodes if episode["season"] == season_number
        )
    return {
        "title": title,
        "seasons": [season_meta[number] for number in sorted(season_meta)],
        "episodes": episodes,
    }


def safe_iframe(source_type: str, source_url: str) -> str:
    if source_type.upper() != "EMBED" or not source_url.startswith(("https://", "http://")):
        return ""
    escaped = html.escape(source_url, quote=True)
    return f'<iframe src="{escaped}" allowfullscreen loading="lazy"></iframe>'


def parse_episode_page(document: str) -> dict[str, Any]:
    root = lxml_html.fromstring(document)
    textareas = root.xpath('//textarea[@id="player-data"]')
    if not textareas:
        return {"status": "unavailable", "error": "player-data ausente", "sources": []}
    payload_text = textareas[0].text or textareas[0].text_content() or ""
    payload = json.loads(payload_text)
    sources: list[dict[str, Any]] = []
    for position, source in enumerate(payload.get("sources") or [], start=1):
        source_type = str(source.get("source_type") or "UNKNOWN").upper()
        source_url = str(source.get("source_url") or "")
        resolved_url = str(source.get("resolved_embed_url") or "")
        effective_url = resolved_url or source_url
        sources.append(
            {
                "source_number": position,
                "source_id": source.get("id"),
                "label": source.get("label") or f"Fonte {position}",
                "type": source_type,
                "audio_type": source.get("audio_type"),
                "audio_language_code": source.get("audio_language_code"),
                "audio_language_label": source.get("audio_language_label"),
                "subtitle_language_code": source.get("subtitle_language_code"),
                "subtitle_language_label": source.get("subtitle_language_label"),
                "source_url": source_url,
                "resolved_url": resolved_url,
                "provider_host": source.get("resolved_provider_host") or urlparse(effective_url).netloc,
                "resolver_status": source.get("embed_resolver_status"),
                "health_status": source.get("health_status"),
                "active": source.get("active"),
                "iframe": safe_iframe(source_type, effective_url),
            }
        )
    return {
        "status": "complete" if sources else "unavailable",
        "error": "" if sources else "nenhuma fonte declarada",
        "sources": sources,
        "meta": payload.get("meta") or {},
        "subtitles": payload.get("subtitles") or [],
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    temp_path = path.with_suffix(path.suffix + ".tmp")
    with temp_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    replace_with_retry(temp_path, path)


def replace_with_retry(source: Path, destination: Path, attempts: int = 8) -> None:
    """Contorna bloqueios transitórios de arquivos no Windows."""
    for attempt in range(attempts):
        try:
            source.replace(destination)
            return
        except PermissionError:
            if attempt + 1 >= attempts:
                raise
            time.sleep(0.25 * (attempt + 1))


def flatten_catalog(catalog: dict[str, Any]) -> tuple[list[dict[str, Any]], ...]:
    anime_rows: list[dict[str, Any]] = []
    season_rows: list[dict[str, Any]] = []
    episode_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    for anime in sorted(catalog.get("animes", []), key=lambda item: item["order"]):
        anime_rows.append(
            {
                "order": anime["order"],
                "title": anime.get("title"),
                "slug": anime.get("slug"),
                "url": anime.get("url"),
                "status": anime.get("status"),
                "season_count": len(anime.get("seasons") or []),
                "episode_count": len(anime.get("episodes") or []),
                "completed_episodes": sum(
                    1 for episode in anime.get("episodes") or [] if episode.get("status") == "complete"
                ),
                "error": anime.get("error", ""),
                "collected_at": anime.get("collected_at", ""),
            }
        )
        for season in anime.get("seasons") or []:
            season_rows.append(
                {
                    "anime_order": anime["order"],
                    "anime": anime.get("title"),
                    "anime_url": anime.get("url"),
                    **season,
                }
            )
        for episode in anime.get("episodes") or []:
            episode_base = {
                "anime_order": anime["order"],
                "anime": anime.get("title"),
                "anime_url": anime.get("url"),
                "season": episode.get("season"),
                "episode": episode.get("episode"),
                "code": episode.get("code"),
                "episode_url": episode.get("url"),
                "status": episode.get("status"),
                "source_count": len(episode.get("sources") or []),
                "error": episode.get("error", ""),
                "collected_at": episode.get("collected_at", ""),
            }
            episode_rows.append(episode_base)
            for source in episode.get("sources") or []:
                source_rows.append({**episode_base, **source})
    return anime_rows, season_rows, episode_rows, source_rows


def save_checkpoint(output_dir: Path, catalog: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    catalog["updated_at"] = utc_now()
    json_path = output_dir / "catalogo.json"
    temp_json = json_path.with_suffix(".json.tmp")
    temp_json.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    replace_with_retry(temp_json, json_path)

    anime_rows, season_rows, episode_rows, source_rows = flatten_catalog(catalog)
    write_csv(
        output_dir / "animes.csv",
        anime_rows,
        [
            "order", "title", "slug", "url", "status", "season_count", "episode_count",
            "completed_episodes", "error", "collected_at",
        ],
    )
    write_csv(
        output_dir / "temporadas.csv",
        season_rows,
        ["anime_order", "anime", "anime_url", "season", "title", "expected_episodes", "listed_episodes"],
    )
    write_csv(
        output_dir / "episodios.csv",
        episode_rows,
        [
            "anime_order", "anime", "anime_url", "season", "episode", "code", "episode_url",
            "status", "source_count", "error", "collected_at",
        ],
    )
    write_csv(
        output_dir / "fontes.csv",
        source_rows,
        [
            "anime_order", "anime", "anime_url", "season", "episode", "code", "episode_url",
            "source_number", "source_id", "label", "type", "audio_type", "audio_language_code",
            "audio_language_label", "subtitle_language_code", "subtitle_language_label", "source_url",
            "resolved_url", "provider_host", "resolver_status", "health_status", "active", "iframe",
            "status", "error", "collected_at",
        ],
    )


def sleep_polite(delay: float) -> None:
    time.sleep(max(0.0, delay) + random.uniform(0.05, 0.25))


def collect_anime(record: dict[str, Any], episode_delay: float) -> dict[str, Any]:
    result = {**record, "status": "in_progress", "error": "", "seasons": [], "episodes": []}
    try:
        page = fetch(record["url"])
        parsed = parse_anime_page(page.url, page.text)
        result["title"] = parsed["title"] or result["title"]
        result["seasons"] = parsed["seasons"]
        result["episodes"] = parsed["episodes"]
        if not result["episodes"]:
            result["status"] = "unavailable"
            result["error"] = "nenhum episodio encontrado na pagina da obra"
            result["collected_at"] = utc_now()
            return result
        failed = 0
        for index, episode in enumerate(result["episodes"]):
            if index:
                sleep_polite(episode_delay)
            try:
                episode_page = fetch(episode["url"])
                episode_result = parse_episode_page(episode_page.text)
                episode.update(episode_result)
                episode["collected_at"] = utc_now()
            except Exception as exc:  # checkpoint keeps the exact failed episode
                failed += 1
                episode["status"] = "error"
                episode["error"] = f"{type(exc).__name__}: {exc}"
                episode["sources"] = []
                episode["collected_at"] = utc_now()
        result["status"] = "complete" if failed == 0 else "partial"
        if failed:
            result["error"] = f"{failed} episodio(s) com erro"
    except Exception as exc:
        result["status"] = "error"
        result["error"] = f"{type(exc).__name__}: {exc}"
    result["collected_at"] = utc_now()
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audita temporadas, episodios e fontes publicas do FIB Anime.")
    parser.add_argument("--input", type=Path, required=True, help="Planilha .xlsx com a aba Animes.")
    parser.add_argument("--output-dir", type=Path, required=True, help="Pasta dos CSVs, JSON e checkpoint.")
    parser.add_argument("--start-order", type=int, default=1, help="Primeira ordem do sitemap a processar.")
    parser.add_argument("--limit-animes", type=int, default=0, help="Limite de obras nesta execucao; 0 = todas.")
    parser.add_argument("--anime-delay", type=float, default=1.0, help="Pausa minima entre obras.")
    parser.add_argument("--episode-delay", type=float, default=0.65, help="Pausa minima entre episodios.")
    parser.add_argument("--force", action="store_true", help="Reprocessa obras ja concluidas.")
    parser.add_argument(
        "--metadata-only",
        action="store_true",
        help="Atualiza apenas titulo e temporadas, preservando episodios e fontes ja coletados.",
    )
    return parser.parse_args()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    args = parse_args()
    inventory = load_inventory(args.input.resolve())
    args.output_dir.mkdir(parents=True, exist_ok=True)
    catalog_path = args.output_dir / "catalogo.json"
    if catalog_path.exists():
        catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    else:
        catalog = {
            "source_workbook": str(args.input.resolve()),
            "site": BASE_URL,
            "created_at": utc_now(),
            "updated_at": utc_now(),
            "animes": [],
        }
    existing = {int(item["order"]): item for item in catalog.get("animes", [])}
    selected = [record for record in inventory if record["order"] >= args.start_order]
    if args.limit_animes > 0:
        selected = selected[: args.limit_animes]

    print(f"Inventario: {len(inventory)} animes; selecionados nesta execucao: {len(selected)}", flush=True)
    if args.metadata_only:
        updated = 0
        for record in selected:
            previous = existing.get(record["order"])
            if not previous or not any(not season.get("title") for season in previous.get("seasons", [])):
                continue
            if updated:
                sleep_polite(args.anime_delay)
            print(f"[{record['order']:03d}] {previous.get('title')}: atualizando metadados", flush=True)
            page = fetch(record["url"])
            parsed = parse_anime_page(page.url, page.text)
            previous["title"] = parsed["title"] or previous.get("title")
            previous["seasons"] = parsed["seasons"]
            previous["metadata_updated_at"] = utc_now()
            existing[record["order"]] = previous
            catalog["animes"] = [existing[key] for key in sorted(existing)]
            save_checkpoint(args.output_dir, catalog)
            updated += 1
        print(f"Metadados atualizados: {updated}", flush=True)
        return 0

    processed_now = 0
    for record in selected:
        previous = existing.get(record["order"])
        if previous and previous.get("status") in {"complete", "unavailable"} and not args.force:
            print(f"[{record['order']:03d}] {previous.get('title')}: ja auditado", flush=True)
            continue
        if processed_now:
            sleep_polite(args.anime_delay)
        print(f"[{record['order']:03d}] {record['title']}: iniciando", flush=True)
        collected = collect_anime(record, args.episode_delay)
        existing[record["order"]] = collected
        catalog["animes"] = [existing[key] for key in sorted(existing)]
        save_checkpoint(args.output_dir, catalog)
        processed_now += 1
        complete = sum(1 for ep in collected.get("episodes", []) if ep.get("status") == "complete")
        total_sources = sum(len(ep.get("sources") or []) for ep in collected.get("episodes", []))
        print(
            f"[{record['order']:03d}] {collected.get('title')}: {collected['status']} | "
            f"{complete}/{len(collected.get('episodes', []))} episodios | {total_sources} fontes",
            flush=True,
        )

    save_checkpoint(args.output_dir, catalog)
    print(f"Saida: {args.output_dir.resolve()}", flush=True)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Interrompido; o ultimo checkpoint concluido foi preservado.", file=sys.stderr)
        raise SystemExit(130)
